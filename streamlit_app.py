import streamlit as st
from openai import OpenAI
import io
from pypdf import PdfReader
from docx import Document

client = OpenAI(api_key=st.secrets["OPENAI_API_KEY"])
def extract_text_from_file(uploaded_file):
    if uploaded_file is None:
        return ""

    file_name = uploaded_file.name.lower()

    try:
        if file_name.endswith(".pdf"):
            uploaded_file.seek(0)
            reader = PdfReader(uploaded_file)

            text = ""
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"

            uploaded_file.seek(0)
            return text

                elif file_name.endswith(".docx"):
            uploaded_file.seek(0)
            document = Document(uploaded_file)

            text_parts = []

            # Read normal paragraphs
            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())

            # Read text inside tables
            for table in document.tables:
                for row in table.rows:
                    row_text = []

                    for cell in row.cells:
                        cell_text = cell.text.strip()

                        if cell_text:
                            row_text.append(cell_text)

                    if row_text:
                        text_parts.append(" | ".join(row_text))

            text = "\n".join(text_parts)

            uploaded_file.seek(0)
            return text

        else:
            return ""

    except Exception:
        uploaded_file.seek(0)
        return ""
        
st.set_page_config(
    page_title="Teacher AI",
    page_icon="📚",
    layout="wide"
)

# ---------- HEADER ----------

st.title("📚 Teacher AI")
st.caption("Your adaptive AI teaching planner")

# ---------- NAVIGATION ----------

page = st.radio(
    "Navigation",
    ["Today", "Teacher Profile", "Planning Setup"],
    horizontal=True,
    label_visibility="collapsed"
)

st.divider()

# ---------- TODAY ----------

if page == "Today":

    st.header("Today's Plan")
    st.write(
        "Your teaching day will appear here based on your timetable, "
        "plans and actual classroom progress."
    )

    if st.button("✨ Generate Today's Plan", type="primary"):
        with st.spinner("Connecting to Teacher AI..."):
            try:
                response = client.responses.create(
                    model="gpt-5.4-mini",
input=(
    "You are Teacher AI, an adaptive planning assistant for primary school teachers.\n\n"
    f"MONTHLY PLAN:\n{st.session_state.get('planning_setup', {}).get('monthly_plan_text', '')}\n\n"
    "This is a monthly-plan reading test. "
    "Using only the monthly plan above, identify: "
    "1. The main Maths topics or objectives. "
    "2. The main English topics or objectives. "
    "3. One Science topic or objective. "
    "Keep the response brief. "
    "Do not invent information that is not contained in the monthly plan. "
    "If no monthly plan is available, reply exactly: "
    "No Monthly Plan has been saved yet."
)
                )

                st.success(response.output_text)

            except Exception as e:
                st.error(f"Connection error: {e}")

    st.subheader("Lessons")

    st.info(
        "No lessons generated yet. Once Teacher AI is connected, "
        "your timetable and lessons will appear here."
    )

# ---------- TEACHER PROFILE ----------

elif page == "Teacher Profile":

    st.header("Teacher Profile")

    st.write(
        "Tell Teacher AI about your class, teaching preferences and "
        "classroom setup. This context will shape future lesson planning."
    )

    # ----- CLASS -----

    st.subheader("Class")

    col1, col2 = st.columns(2)

    with col1:
        class_level = st.selectbox(
            "Class level",
            [
                "Junior Infants",
                "Senior Infants",
                "1st Class",
                "2nd Class",
                "3rd Class",
                "4th Class",
                "5th Class",
                "6th Class"
            ],
            index=6
        )

    with col2:
        language = st.selectbox(
            "School language",
            ["English-medium", "Irish-medium"]
        )

    pupil_count = st.number_input(
        "Number of pupils",
        min_value=1,
        max_value=40,
        value=27
    )

    st.divider()

    # ----- PROGRAMMES -----

    st.subheader("Books & Programmes")

    st.caption(
        "Add the main books or programmes you use. "
        "Teacher AI should never invent specific pages or exercises "
        "unless their contents are available."
    )

    programmes = st.text_area(
        "Books / programmes",
        placeholder=(
            "e.g.\n"
            "Maths — Busy at Maths 5, Work It Out 5\n"
            "English — Starlight 5\n"
            "Gaeilge — Abair Liom 5, Am don Léamh 5\n"
            "SESE — Explore With Me 5"
        ),
        height=140
    )

    st.divider()

    # ----- RESOURCES -----

    st.subheader("Classroom Resources")

    resources = st.text_area(
        "Resources available",
        placeholder=(
            "e.g. interactive touchscreen, mini-whiteboards, "
            "Chromebooks, dice, counters, maths manipulatives, "
            "A4/A3 paper..."
        ),
        height=120
    )

    st.divider()

    # ----- TEACHING STYLE -----

    st.subheader("Teaching Style")

    st.write(
        "Choose approaches you generally enjoy using. "
        "Teacher AI should still select activities based on what "
        "best suits each lesson."
    )

    preferred_methods = st.multiselect(
        "I like using:",
        [
            "Active learning",
            "Mini-games and challenges",
            "Mystery / detective / mission activities",
            "Pupil vs teacher challenges",
            "Mini-whiteboards",
            "Movement",
            "Real-life scenarios",
            "Hands-on activities / manipulatives",
            "Pupil choice",
            "Partner challenges",
            "Creative / designing / making",
            "Prediction / surprise / reveal",
            "Technology / interactive board",
            "Explicit teacher modelling",
            "Independent practice",
            "Discussion"
        ]
    )

    minimise_methods = st.multiselect(
        "I prefer to minimise:",
        [
            "Long teacher explanations",
            "Group work",
            "Pair work",
            "Textbook-heavy lessons",
            "Worksheets",
            "Technology",
            "Competition",
            "Movement",
            "Creative activities"
        ]
    )

    teaching_notes = st.text_area(
        "Anything else about how you like to teach?",
        placeholder=(
            "e.g. Keep lessons practical and engaging. "
            "I like concise plans that I can understand quickly."
        ),
        height=100
    )

    st.divider()

    # ----- DIFFERENTIATION -----

    st.subheader("Differentiation & Additional Needs")

    irish_exemptions = st.number_input(
        "Number of pupils with Irish exemptions",
        min_value=0,
        max_value=40,
        value=0
    )

    differentiation = st.text_area(
        "Differentiation / additional needs",
        placeholder=(
            "Describe recurring differentiation needs or classroom "
            "arrangements Teacher AI should account for."
        ),
        height=120
    )

    st.divider()

    # ----- RECURRING ARRANGEMENTS -----

    st.subheader("Recurring Classroom Arrangements")

    recurring = st.text_area(
        "Regular arrangements Teacher AI should know",
        placeholder=(
            "e.g. fixed PE times, recurring support teaching, "
            "Religion arrangements, Chromebook access, "
            "regular classroom routines..."
        ),
        height=120
    )

    st.divider()

    # ----- REFERENCE MATERIALS -----

    st.subheader("Reference Materials")

    st.caption(
        "Optional examples or templates that can later help Teacher AI "
        "match how you like lessons and resources presented."
    )

    reference_files = st.file_uploader(
        "Upload lesson templates or example resources",
        type=["pdf", "docx", "pptx"],
        accept_multiple_files=True
    )

    st.divider()

    if st.button("Save Teacher Profile", type="primary"):
        st.session_state["teacher_profile"] = {
            "class_level": class_level,
            "language": language,
            "pupil_count": pupil_count,
            "programmes": programmes,
            "resources": resources,
            "preferred_methods": preferred_methods,
            "minimise_methods": minimise_methods,
            "teaching_notes": teaching_notes,
            "irish_exemptions": irish_exemptions,
            "differentiation": differentiation,
            "recurring_arrangements": recurring
        }

        st.success("Teacher Profile saved for this session.")

# ---------- PLANNING SETUP ----------

elif page == "Planning Setup":

    st.header("Planning Setup")

    st.write(
        "Give Teacher AI the information it needs to understand "
        "what you intend to teach and when."
    )

    st.subheader("Weekly Timetable")

    timetable = st.file_uploader(
        "Upload timetable",
        type=["pdf", "docx", "xlsx", "png", "jpg", "jpeg"]
    )

    st.subheader("Monthly Plan")

    monthly_plan = st.file_uploader(
        "Upload current monthly plan",
        type=["pdf", "docx"]
    )

    st.subheader("Yearly Plan")

    yearly_plan = st.file_uploader(
        "Upload yearly plan (optional)",
        type=["pdf", "docx"]
    )

    if st.button("Save Planning Setup", type="primary"):

        timetable_text = extract_text_from_file(timetable)
        monthly_plan_text = extract_text_from_file(monthly_plan)
        yearly_plan_text = extract_text_from_file(yearly_plan)

        st.session_state["planning_setup"] = {
            "timetable_text": timetable_text,
            "monthly_plan_text": monthly_plan_text,
            "yearly_plan_text": yearly_plan_text
        }

        st.success("Planning setup saved and documents processed.")
